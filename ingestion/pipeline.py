import hashlib
import io
from pathlib import Path
from pypdf import PdfReader
from sqlalchemy.ext.asyncio import AsyncSession
from apps.api.app.config import get_settings
from apps.api.app.models import KnowledgeChunk, KnowledgeDocument
from ai.embeddings.providers import HashEmbeddingProvider
from ai.rag.chunking import structure_aware_chunks
from ai.safety.prompt_injection import flag_untrusted_instructions

ALLOWED = {".pdf", ".txt", ".md", ".docx"}


def extract_pages(data: bytes, suffix: str) -> tuple[list[tuple[int, str]], str]:
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        pages = [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
        method = "native_pdf"
        if sum(len(text.strip()) for _, text in pages) < 80:
            raise ValueError("PDF contains insufficient native text; configure the optional OCR runtime")
        return pages, method
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("DOCX extraction requires the optional python-docx package") from exc
        document = Document(io.BytesIO(data))
        return [(1, "\n\n".join(p.text for p in document.paragraphs))], "docx"
    return [(1, data.decode("utf-8"))], "text"


async def ingest(db: AsyncSession, filename: str, data: bytes, metadata: dict) -> KnowledgeDocument:
    settings = get_settings()
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED:
        raise ValueError("Unsupported file type")
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise ValueError("File exceeds upload limit")
    checksum = hashlib.sha256(data).hexdigest()
    document = KnowledgeDocument(title=metadata.get("title") or Path(filename).stem, filename=Path(filename).name, checksum=checksum, source_organization=metadata.get("source_organization"), source_url=metadata.get("source_url"), species=metadata.get("species"), topic=metadata.get("topic"), language=metadata.get("language", "en"), is_trusted=bool(metadata.get("is_trusted")))
    db.add(document)
    await db.flush()
    try:
        pages, method = extract_pages(data, suffix)
        chunks = structure_aware_chunks(pages, settings.rag_chunk_tokens, settings.rag_chunk_overlap)
        embedder = HashEmbeddingProvider(dimension=settings.embedding_dimension)
        vectors = await embedder.embed([chunk.text for chunk in chunks])
        for chunk, vector in zip(chunks, vectors):
            db.add(KnowledgeChunk(document_id=document.id, chunk_number=chunk.number, content=chunk.text, page_number=chunk.page, section_heading=chunk.section, embedding=vector, embedding_model=embedder.model_name, metadata_json={"extraction_method": method, "injection_markers": flag_untrusted_instructions(chunk.text)}))
        document.status = "ready"
    except Exception as exc:
        document.status, document.error = "failed", str(exc)
    await db.commit()
    await db.refresh(document)
    return document
